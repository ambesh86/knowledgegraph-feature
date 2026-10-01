import logging
from pathlib import Path

from pptx import Presentation

from tpp.model.question_type_enum import QuestionTypeEnum
from tpp.model.theraputic_area_enum import TheraputicAreaEnum
from tpp.model.tpp_question import TppQuestion
from tpp.model.tpp import Tpp

logger = logging.getLogger(__name__)


class TppMapper:
    """
    Read, parse and convert a pptx into a collection of Tpp objects
    """

    def map(self, pptx: Path) -> set[Tpp] | None:
        if pptx is None:
            return None

        tpps = self.read_pptx(pptx)
        logger.debug(f"parsed tpps: {tpps}")
        return tpps

    def read_pptx(self, pptx: Path) -> set[Tpp]:
        presentation = Presentation(str(pptx))
        tpps = set()
        for slide in presentation.slides:
            tpp = self._map_slide_to_tpp(slide)
            if tpp is not None:
                tpps.add(tpp)

        for tpp in tpps:
            logger.debug(f"tpp: {tpp}")
        return tpps

    def _map_slide_to_tpp(self, slide) -> Tpp | None:
        if slide is None:
            return None

        is_tpp_slide = False
        product_description = []
        questions = set()
        tpp = None
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                current_shape_content = shape.text.strip()
                if current_shape_content.startswith("Target Product Profile"):
                    logger.info(f"found a target product profile slide")
                    is_tpp_slide = True

                if is_tpp_slide:
                    product_description.append(current_shape_content)

            if shape.has_table:
                for row in shape.table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    # logger.info(f"row: {row_data}")
                    tpp_question = self._map_row_to_question(row_data)
                    if tpp_question is not None:
                        questions.add(tpp_question)

            if is_tpp_slide:
                # todo: map theraputic area
                tpp = Tpp(
                    threaputic_area=TheraputicAreaEnum.UNKNOWN,
                    product_description=self._clean_tpp_product_description(
                        " ".join(product_description)
                    ),
                    questions=questions,
                )

        return tpp

    def _map_row_to_question(self, row: list[str]) -> TppQuestion | None:
        """
        map slide row to a tpp question
        """
        if row is None or len(row) < 4:
            return None

        question_type = self._map_col_to_question_type(row[0])
        logger.debug(f"found question type: {question_type}")
        if question_type is None:
            return None

        return TppQuestion(
            question_type=question_type,
            ideal=row[1],
            acceptable=row[2],
            excluded=row[3],
        )

    def _map_col_to_question_type(self, row_header: str) -> QuestionTypeEnum | None:
        if row_header is None:
            return None
        question_type = row_header.strip().lower()
        lookup = {
            QuestionTypeEnum.INDICATION.value[1]: QuestionTypeEnum.INDICATION,
            QuestionTypeEnum.CONTRAINDICATION.value[
                1
            ]: QuestionTypeEnum.CONTRAINDICATION,
            QuestionTypeEnum.MECHANISM_OF_ACTION.value[
                1
            ]: QuestionTypeEnum.MECHANISM_OF_ACTION,
            QuestionTypeEnum.ROUTE_OF_ADMINISTRATION.value[
                1
            ]: QuestionTypeEnum.ROUTE_OF_ADMINISTRATION,
            QuestionTypeEnum.EFFICACY.value[1]: QuestionTypeEnum.EFFICACY,
            QuestionTypeEnum.SAFTEY_AND_TOLERABILITY.value[
                1
            ]: QuestionTypeEnum.SAFTEY_AND_TOLERABILITY,
            QuestionTypeEnum.COST_OF_GOODS_SOLD.value[
                1
            ]: QuestionTypeEnum.COST_OF_GOODS_SOLD,
        }

        if question_type in lookup:
            return lookup[question_type]
        else:
            return None

    def _clean_tpp_product_description(self, product_description: str) -> str:
        remove_list = [
            "Confidential – Internal Use Only",
            "|",
            "Target Product Profile",
        ]
        for remove in remove_list:
            product_description = self._remove_substring(
                product_description, remove
            ).strip()
        return product_description

    def _remove_substring(self, value: str, remove_term: str) -> str:
        if value.find(remove_term) > -1:
            value = value.replace(remove_term, "")
        return value
